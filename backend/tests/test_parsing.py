"""Slice 2 gate: parse fidelity on 3 known resumes, across txt, PDF and DOCX
input paths. Every asserted value must come verbatim from the fixture."""

import io
from pathlib import Path

from parsing.extract import extract_text_from_docx, extract_text_from_pdf
from parsing.parser import parse_resume

FIXTURES = Path(__file__).parent / "fixtures"


def load(name: str) -> str:
    return (FIXTURES / name).read_text(encoding="utf-8")


# ---------------------------------------------------------------- fidelity


def test_dense_resume_fidelity():
    p = parse_resume(load("resume_dense.txt"))
    assert p["contact"]["name"] == "Priya Sharma"
    assert p["contact"]["email"] == "priya.sharma@example.com"
    assert p["contact"]["phone"] is not None
    assert p["contact"]["linkedin"] == "linkedin.com/in/priyasharma"
    assert p["contact"]["github"] == "github.com/priyasharma"
    assert "payment systems" in p["sections"]["summary"]
    assert "Python" in p["sections"]["skills"]
    assert "Kafka" in p["sections"]["skills"]
    exp = p["sections"]["experience"]
    assert len(exp) == 2
    assert any("FinEdge Payments" in h for h in exp[0]["header"])
    assert exp[0]["dates"] is not None
    assert len(exp[0]["bullets"]) == 3
    assert exp[0]["bullets"][0].startswith("Led migration")
    assert len(exp[1]["bullets"]) == 2
    edu = p["sections"]["education"]
    assert any("NIT Trichy" in h for e in edu for h in e["header"])
    assert len(p["sections"]["certifications"]) == 2
    assert p["flags"]["has_placeholders"] is False


def test_projects_resume_fidelity():
    p = parse_resume(load("resume_projects.txt"))
    assert p["contact"]["name"] == "Arjun Mehta"
    assert p["contact"]["email"] == "arjun.mehta@example.com"
    assert p["contact"]["github"] == "github.com/arjunm"
    assert "side projects" in p["sections"]["summary"]
    assert "TypeScript" in p["sections"]["skills"]
    projects = p["sections"]["projects"]
    assert len(projects) == 2
    assert any("DevDeck" in h for h in projects[0]["header"])
    assert any("3k monthly users" in b for b in projects[0]["bullets"])
    assert len(p["sections"]["experience"]) == 1


def test_sparse_resume_fidelity():
    p = parse_resume(load("resume_sparse.txt"))
    assert p["contact"]["name"] == "Rahul V"
    assert p["contact"]["email"] == "rahul.v@example.com"
    assert p["contact"]["phone"] is None
    assert p["contact"]["linkedin"] is None
    assert p["sections"]["skills"] == ["Java", "HTML"]
    exp = p["sections"]["experience"]
    assert len(exp) == 1
    assert exp[0]["bullets"] == ["Helped with website"]
    # sparse-ness is visible downstream: few bullets, thin summary
    assert p["stats"]["bullet_count"] == 1


# ------------------------------------------------- pdf/docx extraction paths


def _fixture_as_pdf(text: str) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)
    for line in text.splitlines():
        # latin-1-safe: replace bullets for the PDF writer, parser accepts '-'
        pdf.cell(0, 5, line.replace("•", "-").encode("latin-1", "replace").decode("latin-1"), ln=True)
    return bytes(pdf.output())


def _fixture_as_docx(text: str) -> bytes:
    from docx import Document

    doc = Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_pdf_roundtrip_preserves_structure():
    text = load("resume_dense.txt")
    extracted = extract_text_from_pdf(_fixture_as_pdf(text))
    p = parse_resume(extracted)
    assert p["contact"]["email"] == "priya.sharma@example.com"
    assert len(p["sections"]["experience"]) == 2
    assert len(p["sections"]["experience"][0]["bullets"]) == 3


def test_docx_roundtrip_preserves_structure():
    text = load("resume_dense.txt")
    extracted = extract_text_from_docx(_fixture_as_docx(text))
    p = parse_resume(extracted)
    assert p["contact"]["email"] == "priya.sharma@example.com"
    assert len(p["sections"]["experience"]) == 2
    assert p["sections"]["experience"][0]["bullets"][0].startswith("Led migration")


def test_parser_never_invents_content():
    """No-fabrication contract at the parser level: every parsed string is a
    substring of the source text."""
    text = load("resume_dense.txt")
    p = parse_resume(text)

    def all_strings(obj):
        if isinstance(obj, str):
            yield obj
        elif isinstance(obj, dict):
            for v in obj.values():
                yield from all_strings(v)
        elif isinstance(obj, list):
            for v in obj:
                yield from all_strings(v)

    for s in all_strings({"contact": p["contact"], "sections": p["sections"]}):
        for piece in s.split():
            assert piece in text, f"invented token: {piece!r}"


def test_contact_information_heading_not_in_name():
    """'Contact Information' should be recognized as a section heading and
    not bleed into the parsed name."""
    text = "Dev Jain\nContact Information\ndev@example.com | +91 1234567890"
    p = parse_resume(text)
    assert p["contact"]["name"] == "Dev Jain"
    assert "Contact" not in (p["contact"]["name"] or "")


def test_contact_details_heading_not_in_name():
    text = "Jane Doe\nContact Details\njane@example.com"
    p = parse_resume(text)
    assert p["contact"]["name"] == "Jane Doe"


# --- continuation-merge regression guards (from adversarial verification) ---

def test_dateless_header_between_jobs_not_swallowed():
    """CRITICAL: a capitalised company header between two jobs must start its
    own entry, never merge into the prior bullet (which loses a whole job)."""
    text = (
        "John Doe\nj@e.com\n\nExperience\n"
        "Acme Corporation\n"
        "- Built the settlement service that processes millions of\n"
        "transactions each day across regions\n"
        "- Reduced latency\n"
        "Beta Industries\n"
        "- Fixed critical bugs\n"
        "- Wrote regression tests\n"
    )
    exp = parse_resume(text)["sections"]["experience"]
    assert len(exp) == 2
    assert exp[0]["header"][0] == "Acme Corporation"
    assert exp[1]["header"][0] == "Beta Industries"


def test_section_final_wrapped_bullet_no_spurious_entry():
    """A bullet that wraps on the LAST line of a section must merge, not spawn
    a bogus entry or truncate the bullet."""
    text = (
        "Jane\nj@e.com\n\nExperience\n"
        "Senior Engineer | FinEdge\nJan 2020 - Present\n"
        "- Improved throughput significantly\n"
        "- Migrated the legacy monolith into a set of cleanly separated\n"
        "microservices running on Kubernetes\n"
    )
    exp = parse_resume(text)["sections"]["experience"]
    assert len(exp) == 1
    assert "microservices running on Kubernetes" in exp[0]["bullets"][-1]


def test_wrapped_continuation_with_year_not_phantom_dated_entry():
    """A continuation clause mentioning a year must not become a phantom entry
    carrying a fabricated date."""
    text = (
        "Sam\ns@e.com\n\nExperience\n"
        "Engineer | Acme\n2020 - Present\n"
        "- Scaled the platform to absorb a large surge in traffic during\n"
        "the 2021 holiday season without any downtime\n"
        "- Mentored junior engineers\n"
    )
    exp = parse_resume(text)["sections"]["experience"]
    assert len(exp) == 1
    assert exp[0]["dates"] == "2020 - Present"


def test_spaced_dash_continuation_keeps_space():
    """'cost -' + 'saved' must not glue into '-saved'; only true soft word
    breaks ('e-' + 'commerce') join without a space."""
    text = (
        "N\nn@e.com\n\nExperience\nRole | Co\n"
        "- Reduced infrastructure cost -\n"
        "saved the company big money annually\n"
        "- Led a team\n"
    )
    bullets = parse_resume(text)["sections"]["experience"][0]["bullets"]
    assert "cost - saved" in bullets[0]
    assert "-saved" not in bullets[0]


def test_degree_sign_not_treated_as_bullet():
    """normalize_pdf_text must not split '360 degrees' written with a degree
    sign into two bullets."""
    from parsing.extract import normalize_pdf_text
    out = normalize_pdf_text("• Ran 360° feedback cycles across teams")
    # single bullet line, degree sign preserved, not split
    bullet_lines = [l for l in out.splitlines() if l.strip().startswith("-")]
    assert len(bullet_lines) == 1
    assert "360° feedback" in out
