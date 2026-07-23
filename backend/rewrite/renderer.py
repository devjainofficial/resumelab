"""Render composed markdown to preview HTML, PDF, and DOCX.

Per-template renderers live in `rewrite.templates`. Each template produces
a complete HTML+CSS document; WeasyPrint converts that same HTML to PDF so
the download is pixel-identical to the browser preview.

fpdf2 remains as a bare-bones fallback when WeasyPrint's native libraries
are unavailable (Windows dev without GTK). It emits a legible plain PDF.
"""

from __future__ import annotations

import io
import re

try:  # pragma: no cover - environment probe
    from weasyprint import HTML  # type: ignore

    WEASYPRINT_AVAILABLE = True
except Exception:  # pragma: no cover
    WEASYPRINT_AVAILABLE = False

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from rewrite.templates import render_html
from structures.specs import STRUCTURES

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _strip_md(line: str) -> str:
    return BOLD_RE.sub(r"\1", line)


def _strip_watermark(markdown: str) -> str:
    return "\n".join(
        l for l in markdown.splitlines() if not l.startswith("> DRAFT")
    )


def _template_key(structure_id: str) -> str:
    spec = STRUCTURES.get(structure_id, STRUCTURES["S1"])
    return spec.get("template", "classic")


# ------------------------------------------------------------------- HTML


def render_preview_html(markdown: str, structure_id: str) -> str:
    """Full HTML document for the browser iframe preview. Shows the DRAFT
    watermark if the markdown carries one."""
    watermark = any(
        l.startswith("> DRAFT") for l in markdown.splitlines()
    )
    return render_html(_template_key(structure_id), markdown, watermark=watermark)


# -------------------------------------------------------------------- PDF


def _render_pdf_weasyprint(markdown: str, structure_id: str) -> bytes:  # pragma: no cover
    # Watermark NEVER in downloads — only preview.
    html_doc = render_html(_template_key(structure_id), markdown, watermark=False)
    return HTML(string=html_doc).write_pdf()


def _render_pdf_fpdf(markdown: str) -> bytes:
    """Fallback used only when WeasyPrint isn't available (Windows dev).
    Produces a legible-but-plain PDF regardless of template."""
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF(format="A4")
    pdf.set_margins(14, 12, 14)
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()

    def safe(s: str) -> str:
        return s.encode("latin-1", "replace").decode("latin-1")

    def block(text: str, height: float) -> None:
        pdf.multi_cell(0, height, text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    prev_was_name = False
    for raw in _strip_watermark(markdown).splitlines():
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            pdf.set_font("Helvetica", "B", 18)
            pdf.set_text_color(10, 10, 10)
            block(safe(line[2:]), 8)
            prev_was_name = True
            continue
        elif line.startswith("## "):
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 10)
            pdf.set_text_color(34, 34, 34)
            block(safe(line[3:].upper()), 5)
            y = pdf.get_y()
            pdf.set_draw_color(200, 200, 200)
            pdf.line(14, y, 196, y)
            pdf.ln(1.5)
        elif prev_was_name and "|" in line:
            pdf.set_font("Helvetica", "", 9)
            pdf.set_text_color(85, 85, 85)
            block(safe(_strip_md(line)), 4)
            y = pdf.get_y() + 1
            pdf.set_draw_color(34, 34, 34)
            pdf.set_line_width(0.5)
            pdf.line(14, y, 196, y)
            pdf.set_line_width(0.2)
            pdf.ln(3)
        elif line.startswith("**"):
            pdf.ln(1)
            pdf.set_font("Helvetica", "B", 10)
            pdf.set_text_color(26, 26, 26)
            block(safe(_strip_md(line)), 4.5)
        elif line.startswith("- "):
            pdf.set_font("Helvetica", "", 9.5)
            pdf.set_text_color(34, 34, 34)
            text = safe(_strip_md(line[2:]))
            pdf.cell(4, 4.5, "-", new_x=XPos.RIGHT, new_y=YPos.TOP)
            pdf.multi_cell(0, 4.5, text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        else:
            pdf.set_font("Helvetica", "", 9.5)
            pdf.set_text_color(51, 51, 51)
            block(safe(_strip_md(line)), 4.5)
        prev_was_name = False

    return bytes(pdf.output())


def render_pdf(markdown: str, structure_id: str = "S1") -> bytes:
    if WEASYPRINT_AVAILABLE:  # pragma: no cover
        return _render_pdf_weasyprint(markdown, structure_id)
    return _render_pdf_fpdf(markdown)


# ------------------------------------------------------------------ DOCX


def render_docx(markdown: str) -> bytes:
    """Structural DOCX — deliberately simpler than the visual templates so
    ATS parsers reliably extract text. All templates share this DOCX."""
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    style.paragraph_format.space_after = Pt(1)
    style.paragraph_format.space_before = Pt(0)

    prev_was_name = False
    for raw in _strip_watermark(markdown).splitlines():
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(line[2:])
            run.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0x0A, 0x0A, 0x0A)
            p.paragraph_format.space_after = Pt(0)
            prev_was_name = True
            continue
        elif line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line[3:].upper())
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
            from docx.oxml.ns import qn
            pPr = p._element.get_or_add_pPr()
            pBdr = pPr.makeelement(qn("w:pBdr"), {})
            bottom = pBdr.makeelement(
                qn("w:bottom"),
                {
                    qn("w:val"): "single",
                    qn("w:sz"): "4",
                    qn("w:space"): "1",
                    qn("w:color"): "CCCCCC",
                },
            )
            pBdr.append(bottom)
            pPr.append(pBdr)
        elif prev_was_name and "|" in line:
            p = doc.add_paragraph()
            run = p.add_run(_strip_md(line))
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            p.paragraph_format.space_after = Pt(4)
        elif line.startswith("**"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(1)
            cleaned = _strip_md(line)
            parts = cleaned.split(" — ", 1)
            run = p.add_run(parts[0])
            run.bold = True
            run.font.size = Pt(10)
            if len(parts) > 1:
                run2 = p.add_run(f" — {parts[1]}")
                run2.font.size = Pt(9)
                run2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif line.startswith("- "):
            p = doc.add_paragraph(_strip_md(line[2:]), style="List Bullet")
            p.paragraph_format.space_after = Pt(0.5)
            for run in p.runs:
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        else:
            p = doc.add_paragraph()
            run = p.add_run(_strip_md(line))
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        prev_was_name = False

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
